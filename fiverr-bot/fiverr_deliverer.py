"""
Fiverr browser automation — Playwright.

Handles:
1. Login to Fiverr
2. Navigating to an order
3. Sending a message (clarifying question, reactions)
4. Delivering the completed work (uploading .docx + message)

Install Playwright browsers once with:
    playwright install chromium
"""

import asyncio
import io
import logging
import os
import random
import tempfile

from docx import Document
from playwright.async_api import async_playwright, Page, BrowserContext

import config
from humanizer import maybe_add_typo, correction_message

log = logging.getLogger(__name__)

FIVERR_BASE = "https://www.fiverr.com"
LOGIN_URL = f"{FIVERR_BASE}/login"


class FiverrClient:
    def __init__(self):
        self._playwright = None
        self._browser = None
        self._context: BrowserContext | None = None
        self._page: Page | None = None
        self._logged_in = False

    async def start(self) -> None:
        self._playwright = await async_playwright().start()
        self._browser = await self._playwright.chromium.launch(
            headless=True,
            args=["--no-sandbox", "--disable-dev-shm-usage"],
        )
        self._context = await self._browser.new_context(
            viewport={"width": 1280, "height": 800},
            user_agent=(
                "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
                "AppleWebKit/537.36 (KHTML, like Gecko) "
                "Chrome/124.0.0.0 Safari/537.36"
            ),
            locale="en-US",
            timezone_id="Europe/Madrid",
        )
        self._page = await self._context.new_page()

    async def stop(self) -> None:
        if self._context:
            await self._context.close()
        if self._browser:
            await self._browser.close()
        if self._playwright:
            await self._playwright.stop()

    async def login(self) -> None:
        if self._logged_in:
            return
        log.info("Logging into Fiverr as %s.", config.FIVERR_EMAIL)
        page = self._page

        await page.goto(LOGIN_URL, wait_until="networkidle")
        await _human_pause(1, 2)

        # Fill email
        await page.fill('input[name="email"], input[type="email"]', config.FIVERR_EMAIL)
        await _human_pause(0.5, 1.2)

        # Fill password
        await page.fill('input[name="password"], input[type="password"]', config.FIVERR_PASSWORD)
        await _human_pause(0.5, 1.5)

        # Submit
        await page.press('input[name="password"], input[type="password"]', "Enter")
        await page.wait_for_load_state("networkidle")

        # Check if login succeeded
        if "login" in page.url or "captcha" in page.url.lower():
            raise RuntimeError(f"Fiverr login failed. Current URL: {page.url}")

        self._logged_in = True
        log.info("Fiverr login successful.")

    async def go_to_order(self, order_id: str) -> bool:
        """Navigate to the order page. Returns True if found."""
        page = self._page
        order_url = f"{FIVERR_BASE}/orders/{order_id}"
        await page.goto(order_url, wait_until="networkidle")
        await _human_pause(1, 2)

        if "404" in await page.title() or "not found" in (await page.title()).lower():
            log.warning("Order %s not found on Fiverr.", order_id)
            return False
        return True

    async def send_message(self, order_id: str, message: str) -> bool:
        """
        Send a message on the order page. Handles optional typo + self-correction.
        Returns True on success.
        """
        page = self._page
        if not await self.go_to_order(order_id):
            return False

        typo_text, has_typo = maybe_add_typo(message)

        success = await _type_and_send(page, typo_text)
        if not success:
            return False

        if has_typo:
            await _human_pause(3, 7)  # short pause before correction
            correction = correction_message(typo_text)
            await _type_and_send(page, correction)

        log.info("Message sent on order %s.", order_id)
        return True

    async def deliver_order(
        self,
        order_id: str,
        client_message: str,
        script_text: str,
        buyer_name: str,
    ) -> bool:
        """
        Deliver the completed order on Fiverr:
        1. Navigate to order
        2. Click Deliver Now
        3. Type the client message
        4. Attach the script as a .docx file
        5. Submit delivery
        """
        page = self._page
        if not await self.go_to_order(order_id):
            return False

        # Create a .docx file from the script
        docx_path = _create_docx(script_text, buyer_name, order_id)

        try:
            # Click "Deliver Now" button
            await page.click(
                'button:has-text("Deliver Now"), a:has-text("Deliver Now")',
                timeout=10000,
            )
            await _human_pause(1, 2)

            # Fill in delivery message
            delivery_textarea = page.locator(
                'textarea[placeholder*="message"], textarea[name="description"], [data-testid="delivery-message"]'
            ).first
            await delivery_textarea.click()
            await _human_pause(0.3, 0.8)
            await _type_slowly(page, delivery_textarea, client_message)
            await _human_pause(0.5, 1.5)

            # Upload file
            file_input = page.locator('input[type="file"]').first
            await file_input.set_input_files(docx_path)
            await _human_pause(2, 4)

            # Submit delivery
            await page.click(
                'button[type="submit"]:has-text("Submit"), button:has-text("Deliver Order")',
                timeout=10000,
            )
            await page.wait_for_load_state("networkidle")

            log.info("Order %s delivered successfully.", order_id)
            return True

        except Exception as exc:
            log.error("Delivery failed for order %s: %s", order_id, exc)
            return False
        finally:
            if os.path.exists(docx_path):
                os.unlink(docx_path)

    async def react_to_message(self, order_id: str) -> None:
        """Add a thumbs up or heart reaction to the latest buyer message."""
        page = self._page
        if not await self.go_to_order(order_id):
            return
        try:
            # Hover over last buyer message to reveal reaction button
            messages = page.locator('[data-testid="message-bubble"], .message-content')
            count = await messages.count()
            if count == 0:
                return
            last_msg = messages.nth(count - 1)
            await last_msg.hover()
            await _human_pause(0.5, 1)

            # Click emoji/reaction button
            await page.click('[aria-label="Add reaction"], [data-testid="emoji-reaction"]', timeout=5000)
            await _human_pause(0.3, 0.8)

            # Pick thumbs up or heart randomly
            emoji = "👍" if random.random() > 0.5 else "❤️"
            await page.click(f'[aria-label="{emoji}"], button:has-text("{emoji}")', timeout=5000)
            log.info("Reacted with %s on order %s.", emoji, order_id)
        except Exception as exc:
            log.debug("Could not add reaction on order %s: %s", order_id, exc)


# ── Helpers ──────────────────────────────────────────────────────────────────

async def _human_pause(low: float, high: float) -> None:
    await asyncio.sleep(random.uniform(low, high))


async def _type_slowly(page: Page, locator, text: str) -> None:
    """Type text with random per-character delays to look human."""
    await locator.fill("")
    for char in text:
        await locator.type(char, delay=random.randint(40, 130))


async def _type_and_send(page: Page, text: str) -> bool:
    """Find the message input, type text, and send."""
    try:
        textarea = page.locator(
            'textarea[placeholder*="message"], '
            '[contenteditable="true"][data-testid*="message"], '
            'textarea[name="message"]'
        ).first
        await textarea.click()
        await _human_pause(0.4, 1.0)
        await _type_slowly(page, textarea, text)
        await _human_pause(0.5, 1.2)
        await page.keyboard.press("Enter")
        return True
    except Exception as exc:
        log.error("Could not send message: %s", exc)
        return False


def _create_docx(script_text: str, buyer_name: str, order_id: str) -> str:
    """Create a Word document from the script and return the temp file path."""
    doc = Document()
    doc.add_heading(f"Video Script — {buyer_name}", 0)
    doc.add_paragraph(f"Order: {order_id}")
    doc.add_paragraph("")  # spacer

    for line in script_text.split("\n"):
        stripped = line.strip()
        if not stripped:
            doc.add_paragraph("")
        elif stripped.startswith("[") and stripped.endswith("]"):
            # Stage directions / recording notes — italic
            p = doc.add_paragraph()
            run = p.add_run(stripped)
            run.italic = True
        else:
            doc.add_paragraph(stripped)

    tmp = tempfile.NamedTemporaryFile(
        suffix=".docx", delete=False, prefix=f"script_{order_id}_"
    )
    doc.save(tmp.name)
    tmp.close()
    return tmp.name
